from pathlib import Path
import pandas as pd,numpy as np,json
from pptx import Presentation
from pptx.oxml.xmlchemy import OxmlElement
from copy import deepcopy
R=Path.cwd();S=R/'data/fig1_sources';D=R/'generated/fig1_curves';D.mkdir(parents=True,exist_ok=True)
source=R/'artwork/Fig1_author_input.pptx'
p=Presentation(source);slide=p.slides[0];lookup={s.name:s for s in slide.shapes}
i=pd.read_csv(S/'historical_operating_inputs_10days.csv');f=pd.read_csv(S/'observed_and_predicted_prices_10days.csv');n=pd.read_csv(S/'network_history_14days.csv');b=pd.read_csv(S/'budget_selection_80days.csv');t=pd.read_csv(S/'experiment_30days.csv');m=pd.read_csv(S/'measured_prices_30days.csv')
series={}
for label,col in [('DA price','price_day_ahead'),('Net demand','Net demand'),('Total supply','elec_gene_total_pred'),('Residual bidding space','Residual bidding space'),('Renewable share','Renewable penetration'),('Conventional generation','Conventional generation'),('PV forecast','PV forecast'),('Wind forecast','Wind forecast')]:series[label+' | recorded samples']=(i[col].to_numpy(),None,None)
for name,arr,ax,allv in [
 ('Historical network net demand | network recorded samples',n.iloc[:,1].to_numpy(),'Historical network net demand axes',n.iloc[:,1].to_numpy()),
 ('Observed and forecast prices | Observed recorded samples',f.observed_DA_CNY_MWh.to_numpy(),'Observed and forecast prices axes',np.r_[f.observed_DA_CNY_MWh,f.LoRA_forecast_CNY_MWh]),
 ('Observed and forecast prices | LoRA forecast recorded samples',f.LoRA_forecast_CNY_MWh.to_numpy(),'Observed and forecast prices axes',np.r_[f.observed_DA_CNY_MWh,f.LoRA_forecast_CNY_MWh]),
 ('Selected energy budget | budget recorded samples',b.budget_each_MWh.to_numpy(),'Selected energy budget axes',b.budget_each_MWh.to_numpy()),
 ('Executed PCC import | PCC recorded samples',t.pcc_mw.to_numpy(),'Executed PCC import axes',np.r_[t.pcc_mw,17.8131]),
 ('Observed settlement prices | DA recorded samples',m.observed_DA_CNY_MWh.to_numpy(),'Observed settlement prices axes',np.r_[m.observed_DA_CNY_MWh,m.observed_RT_CNY_MWh]),
 ('Observed settlement prices | RT recorded samples',m.observed_RT_CNY_MWh.to_numpy(),'Observed settlement prices axes',np.r_[m.observed_DA_CNY_MWh,m.observed_RT_CNY_MWh])]:series[name]=(arr,ax,allv)
report=[];records=[];Q=1000000000
for name,(v,axis,allv) in series.items():
 sh=lookup[name];original_line=sh._element.spPr.find('{http://schemas.openxmlformats.org/drawingml/2006/main}ln').xml
 if axis:
  a=lookup[axis];left,top,width,height=a.left,a.top,a.width,a.height
  lo=float(np.min(allv));hi=float(np.max(allv));pad=max((hi-lo)*.08,.01);lo-=pad;hi+=pad
 else:
  left,top,width,height=sh.left,sh.top,sh.width,sh.height;lo=float(v.min());hi=float(v.max())
 xx=np.linspace(0,Q,len(v));yy=(hi-v)/(hi-lo)*Q
 geom=sh._element.spPr.find('{http://schemas.openxmlformats.org/drawingml/2006/main}custGeom')
 path=geom.xpath('./a:pathLst/a:path')[0]
 oldpts=np.array([[int(pt.get('x')),int(pt.get('y'))] for pt in path.xpath('./a:moveTo/a:pt | ./a:lnTo/a:pt')])
 oldx=sh.left+oldpts[:,0]/int(path.get('w'))*sh.width;oldy=sh.top+oldpts[:,1]/int(path.get('h'))*sh.height
 newx=left+xx/Q*width;newy=top+yy/Q*height
 drift=float(np.max(np.hypot(oldx-newx,oldy-newy))/12700)
 for c in list(path):path.remove(c)
 path.set('w',str(Q));path.set('h',str(Q));path.set('fill','none')
 for j,(x,y) in enumerate(zip(xx,yy)):
  cmd=OxmlElement('a:moveTo' if j==0 else 'a:lnTo');pt=OxmlElement('a:pt');pt.set('x',str(round(x)));pt.set('y',str(round(y)));cmd.append(pt);path.append(cmd)
 sh.left=left;sh.top=top;sh.width=width;sh.height=height
 assert sh._element.spPr.find('{http://schemas.openxmlformats.org/drawingml/2006/main}ln').xml==original_line
 recovered=hi-np.rint(yy)/Q*(hi-lo)
 assert np.max(np.abs(recovered-v))<max((hi-lo)*1e-8,1e-7)
 report.append({'curve':name,'samples':len(v),'range':[float(v.min()),float(v.max())],'line_width_pt':sh.line.width/12700,'maximum_old_coordinate_error_pt':drift,'old_unique_x_positions':len(np.unique(oldpts[:,0])),'new_unique_x_positions':len(np.unique(np.rint(xx)))})
 for j,value in enumerate(v):records.append({'curve':name,'sample':j,'value':value})
p.save(D/'Fig1_user_layout_curves_corrected.pptx')
old=Presentation(source)
for a,c in zip(old.slides[0].shapes,slide.shapes):
 if a.name not in series:assert a._element.xml==c._element.xml,a.name
assert len(slide.shapes)==len(old.slides[0].shapes)
pd.DataFrame(records).to_csv(D/'verified_curve_values.csv',index=False)
(D/'curve_verification.json').write_text(json.dumps(report,indent=2),encoding='utf-8')
(D/'修正说明.txt').write_text('只修正15条曲线的路径坐标。所有非曲线元素逐项XML比对保持不变，曲线颜色、实虚线和线宽原样保留。原曲线仍对应旧版数据路径；主要问题是旧绘制程序将密集采样点取整到较粗坐标网格。本次从保存的原始实验CSV重新生成十亿级坐标精度路径，未平滑、删点或重新选择日期。Stage D 的PCC包含2880个15分钟记录，DA与RT价格各1440个半小时记录。网络运行曲线为SimBench实验输出，不是现场测量。保留原文件，不覆盖。',encoding='utf-8')
print('Updated',len(series),'curves;',len(records),'samples; all other shapes unchanged')
print(json.dumps(report[-3:],indent=2))
